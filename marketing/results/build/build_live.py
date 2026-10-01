"""The live results sheet: a template that becomes a Google Sheet on upload.

Its Data tab pulls https://theoverlay.com.au/api/results.csv with IMPORTDATA,
which Google refreshes about hourly, and every other tab is formulas over it,
so nothing here holds a single result. QUERY, SORT and UNIQUE are Google
Sheets functions: this file is meant for Sheets, not Excel.

    python build_live.py   ->   ../Overlay results (live).xlsx
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L
from PIL import Image

CSV = 'https://theoverlay.com.au/api/results.csv'

# The site's own palette (src/app/globals.css) and its two faces.
INK, PAPER, LIME, LIME_SOFT, GREEN = '14161A', 'F5F7F2', 'C6F24E', 'EDF9C8', '6F9A12'
BLUE, BLUE_SOFT, WAY_SOFT, RED, RED_SOFT = '1F6FD6', 'DCEBFF', 'EEF5FF', 'D93636', 'FFE1E1'
ALT, MUTED, SOFT = 'F7F8F5', '6B716A', 'A9AEA4'
SANS, MONO = 'Archivo', 'IBM Plex Mono'

fill = lambda c: PatternFill('solid', fgColor=c, bgColor=c)
font = lambda face=SANS, **k: Font(name=face, **k)
U = '+0.00;-0.00;0.00'
ROWS = 3000   # room for the call sheets to grow into
TOP = 5       # header row; rows 1 to 4 carry the brand band

# A small logo keeps the template small enough to upload.
logo = Image.open('logo.png').convert('RGBA').resize((210, 53), Image.LANCZOS)
logo.quantize(colors=64, method=Image.FASTOCTREE).save('logo-small.png', optimize=True)


def band(ws, ncols, title, sub):
    for r in range(1, 5):
        ws.row_dimensions[r].height = 22 if r in (1, 4) else 30
        for c in range(1, ncols + 1):
            ws.cell(r, c).fill = fill(INK)
    img = XLImage('logo-small.png'); px = 9525
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=16 * px, row=0, rowOff=42 * px), ext=XDRPositiveSize2D(210 * px, 53 * px))
    ws.add_image(img)
    t = ws.cell(2, ncols, title); t.font = font(bold=True, size=16, color=PAPER); t.alignment = Alignment(horizontal='right', vertical='bottom')
    s = ws.cell(3, ncols, sub); s.font = font(MONO, size=9, color=LIME); s.alignment = Alignment(horizontal='right', vertical='top')
    ws.sheet_view.showGridLines = False


def header(ws, r, labels, left=()):
    for j, h in enumerate(labels, 1):
        c = ws.cell(r, j, h); c.font = font(bold=True, color=LIME, size=10); c.fill = fill(INK)
        c.alignment = Alignment(horizontal='left' if j in left else 'center', vertical='center')
    ws.row_dimensions[r].height = 24


def text_rule(ws, rng, value, color, bg=None, bold=True):
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{value}"'], font=Font(color=color, bold=bold), fill=fill(bg) if bg else None))


def sign_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0.0001'], font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['-0.0001'], font=Font(color=RED, bold=True)))


COLS = ['Date', 'Track', 'Race', 'No.', 'Horse', 'Call', 'Tag', 'Rated $', 'Price $', 'Edge', 'Stake (u)', 'Finish', 'Result', 'Profit (u)', 'At risk (u)']
W = [15, 20, 8, 7, 24, 8, 10, 11, 11, 10, 11, 11, 10, 13, 13]
FMT = {'A': 'd mmm yyyy', 'H': '$0.00', 'I': '$0.00', 'J': '0.0%', 'K': '0.0#', 'N': U, 'O': '0.00'}
LEFT = (2, 5)


def call_sheet(ws, title, sub, where):
    band(ws, len(COLS), title, sub)
    header(ws, TOP, COLS, left=LEFT)
    for i, w in enumerate(W, 1):
        col = L(i); d = ws.column_dimensions[col]; d.width = w
        d.font = font(MONO if col in 'ACDHIJKLNO' else SANS, color=INK, bold=col in 'EN')
        d.alignment = Alignment(horizontal='left' if i in LEFT else 'center', vertical='center')
        if col in FMT: d.number_format = FMT[col]
    q = f'select * where Col1 is not null{where} order by Col1 desc, Col2, Col3'
    ws.cell(TOP + 1, 1, f'=IFERROR(QUERY(Data!A2:O20000,"{q}",0),"Loading results")')
    body = f'A{TOP+1}:O{ROWS}'
    ws.conditional_formatting.add(body, FormulaRule(formula=[f'AND($A{TOP+1}<>"",ISEVEN(ROW()))'], fill=fill(ALT)))
    text_rule(ws, f'G{TOP+1}:G{ROWS}', 'Prime', INK, LIME)
    text_rule(ws, f'G{TOP+1}:G{ROWS}', 'Bet', BLUE, BLUE_SOFT)
    text_rule(ws, f'G{TOP+1}:G{ROWS}', 'Way', BLUE, WAY_SOFT)
    text_rule(ws, f'G{TOP+1}:G{ROWS}', 'Lay', RED, RED_SOFT)
    text_rule(ws, f'F{TOP+1}:F{ROWS}', 'Bet', BLUE)
    text_rule(ws, f'F{TOP+1}:F{ROWS}', 'Lay', RED)
    text_rule(ws, f'M{TOP+1}:M{ROWS}', 'Win', GREEN)
    text_rule(ws, f'M{TOP+1}:M{ROWS}', 'Loss', RED)
    sign_rules(ws, f'N{TOP+1}:N{ROWS}')
    ws.freeze_panes = f'A{TOP+1}'


wb = Workbook()
ov = wb.active; ov.title = 'Overview'
allc = wb.create_sheet('All Calls'); bets = wb.create_sheet('Bets'); lays = wb.create_sheet('Lays'); data = wb.create_sheet('Data')
for ws, c in ((ov, LIME), (allc, INK), (bets, BLUE), (lays, RED), (data, SOFT)): ws.sheet_properties.tabColor = c

SINCE = '="Since "&TEXT(MIN(Data!A2:A20000),"d mmm yyyy")&". Updates hourly."'
call_sheet(allc, 'All calls', SINCE, '')
call_sheet(bets, 'Bets', SINCE, " and Col6 = 'Bet'")
call_sheet(lays, 'Lays', SINCE, " and Col6 = 'Lay'")

data['A1'] = f'=IMPORTDATA("{CSV}")'
data.column_dimensions['A'].width = 12
data['Q1'] = 'Pulled from theoverlay.com.au/api/results.csv, refreshed by Google about hourly. Do not type in this tab.'
data['Q1'].font = font(size=9, color=MUTED)

# ---- Overview: everything is a formula over Data ----
# Fixed ranges: the open-ended Data!N2:N is Sheets-only and breaks when the xlsx converts.
D = lambda c: f'Data!${c}$2:${c}$20000'
NC = 10
band(ov, NC, 'Results', SINCE)
ov.column_dimensions['A'].width = 20
for c in 'BCDEFGHIJ': ov.column_dimensions[c].width = 13

r0 = 6
for label, side, col in (('ALL CALLS', '', 'B'), ('BETS', 'Bet', 'E'), ('LAYS', 'Lay', 'H')):
    ci = ord(col) - 64
    crit = f',{D("F")},"{side}"' if side else ''
    dark = not side
    for rr in range(r0, r0 + 4):
        for cc in range(ci, ci + 3): ov.cell(rr, cc).fill = fill(INK if dark else ALT)
    c = ov.cell(r0, ci, label); c.font = font(bold=True, size=9, color=LIME if dark else MUTED)
    p = ov.cell(r0 + 1, ci, f'=SUMIFS({D("N")}{crit})' if side else f'=SUM({D("N")})')
    p.number_format = '+0.0"u";-0.0"u";0.0"u"'; p.font = font(bold=True, size=24, color=LIME if dark else INK)
    ov.merge_cells(start_row=r0 + 1, start_column=ci, end_row=r0 + 2, end_column=ci + 2)
    p.alignment = Alignment(horizontal='left', vertical='center')
    won = f'COUNTIFS({D("M")},"Win"{crit})'; lost = f'COUNTIFS({D("M")},"Loss"{crit})'
    risk = f'SUMIFS({D("O")}{crit})' if side else f'SUM({D("O")})'
    s = ov.cell(r0 + 3, ci, f'=IFERROR({won}+{lost}&" settled, "&TEXT({won}/({won}+{lost}),"0%")&" strike, "&TEXT({p.coordinate}/{risk},"+0.0%;-0.0%")&" ROI","")')
    s.font = font(MONO, size=8, color=SOFT if dark else MUTED)
    if side: ov.conditional_formatting.add(p.coordinate, CellIsRule(operator='lessThan', formula=['0'], font=Font(color=RED, bold=True)))
for rr in range(r0, r0 + 4): ov.row_dimensions[rr].height = 20

ov.cell(11, 1, 'Level stakes: one unit a call, a tenth on a Way. Bets settle at the best of fixed odds, SP and BSP; lays at the shortest lay price or BSP.').font = font(size=9, color=MUTED)
ov.cell(12, 1, 'A lay wins one unit when the horse loses and risks the price less one. ROI is profit over units at risk. Calls on scratched horses and abandoned races are void and left out.').font = font(size=9, color=MUTED)

HDR = ['', 'Calls', 'Won', 'Lost', 'Strike', 'Profit (u)', 'At risk (u)', 'ROI']


def section(r, title):
    ov.cell(r, 1, title.upper()).font = font(bold=True, size=11, color=GREEN)
    ov.row_dimensions[r].height = 22


def stat_row(r, crit, k, guard=None):
    cs = ''.join(f',{D(c)},{v}' for c, v in crit)
    g = (lambda f: f'=IF({guard},"",{f})') if guard else (lambda f: '=' + f)
    ov.cell(r, 2, g(f'COUNTIFS({D("M")},"<>"{cs})'))
    for j, res in zip((3, 4), ('Win', 'Loss')):
        ov.cell(r, j, g(f'COUNTIFS({D("M")},"{res}"{cs})'))
    ov.cell(r, 5, g(f'IF(C{r}+D{r}=0,"",C{r}/(C{r}+D{r}))')).number_format = '0.0%'
    ov.cell(r, 6, g(f'SUMIFS({D("N")}{cs})' if crit else f'SUM({D("N")})')).number_format = U
    ov.cell(r, 7, g(f'SUMIFS({D("O")}{cs})' if crit else f'SUM({D("O")})')).number_format = '0.00'
    ov.cell(r, 8, g(f'IF(G{r}=0,"",F{r}/G{r})')).number_format = '+0.0%;-0.0%;0.0%'
    for j in range(1, len(HDR) + 1):
        c = ov.cell(r, j); c.font = font(SANS if j == 1 else MONO, bold=j in (1, 6), color=INK)
        c.alignment = Alignment(horizontal='left' if j == 1 else 'center', vertical='center')
        if k % 2: c.fill = fill(ALT)
    ov.row_dimensions[r].height = 20


r = 14
section(r, 'By call'); r += 1
header(ov, r, HDR); r += 1; s = r
for k, (lab, crit) in enumerate((('All calls', []), ('Bets', [('F', '"Bet"')]), ('Lays', [('F', '"Lay"')]))):
    ov.cell(r, 1, lab); stat_row(r, crit, k); r += 1
sign_rules(ov, f'F{s}:F{r-1}'); sign_rules(ov, f'H{s}:H{r-1}')
r += 1

section(r, 'By tag'); r += 1
header(ov, r, ['Tag'] + HDR[1:], left=(1,)); r += 1; s = r
for k, (t, bg, fg) in enumerate((('Prime', LIME, INK), ('Bet', BLUE_SOFT, BLUE), ('Way', WAY_SOFT, BLUE), ('Lay', RED_SOFT, RED))):
    ov.cell(r, 1, t); stat_row(r, [('G', f'"{t}"')], k)
    ov.cell(r, 1).fill = fill(bg); ov.cell(r, 1).font = font(bold=True, color=fg)
    r += 1
sign_rules(ov, f'F{s}:F{r-1}'); sign_rules(ov, f'H{s}:H{r-1}')
r += 1

# The last twelve weeks, newest first, Monday to Sunday.
section(r, 'Last 12 weeks'); r += 1
header(ov, r, ['Week of'] + HDR[1:], left=(1,)); r += 1; s = r
for k in range(12):
    a = ov.cell(r, 1, f'=IFERROR(MAX({D("A")})-WEEKDAY(MAX({D("A")}),3),"")' if k == 0 else f'=IF(A{r-1}="","",IF(A{r-1}-7<MIN({D("A")})-6,"",A{r-1}-7))')
    stat_row(r, [('A', f'">="&$A{r}'), ('A', f'"<"&($A{r}+7)')], k, guard=f'$A{r}=""')
    a.number_format = 'd mmm yyyy'
    r += 1
sign_rules(ov, f'F{s}:F{r-1}'); sign_rules(ov, f'H{s}:H{r-1}')

# Day by day, oldest first, feeding the chart: columns L to P.
dr = 30
ov.cell(dr - 1, 12, 'DAY BY DAY').font = font(bold=True, size=11, color=GREEN)
for j, h in enumerate(['Date', 'Bets (u)', 'Lays (u)', 'Day (u)', 'Running (u)'], 12):
    c = ov.cell(dr, j, h); c.font = font(bold=True, color=LIME, size=10); c.fill = fill(INK); c.alignment = Alignment(horizontal='center')
for col, w in zip('LMNOP', (14, 11, 11, 11, 13)): ov.column_dimensions[col].width = w
ov.cell(dr + 1, 12, f'=IFERROR(SORT(UNIQUE(FILTER({D("A")},{D("A")}<>""))),"")')
DAYS = 400
for i in range(dr + 1, dr + 1 + DAYS):
    ov.cell(i, 12).number_format = 'ddd d mmm'
    ov.cell(i, 13, f'=IF($L{i}="","",SUMIFS({D("N")},{D("A")},$L{i},{D("F")},"Bet"))').number_format = U
    ov.cell(i, 14, f'=IF($L{i}="","",SUMIFS({D("N")},{D("A")},$L{i},{D("F")},"Lay"))').number_format = U
    ov.cell(i, 15, f'=IF($L{i}="","",M{i}+N{i})').number_format = U
    ov.cell(i, 16, f'=IF($L{i}="","",SUMIFS({D("N")},{D("A")},"<="&$L{i}))').number_format = U
    for j in range(12, 17):
        ov.cell(i, j).font = font(MONO, size=9, color=INK); ov.cell(i, j).alignment = Alignment(horizontal='center')
sign_rules(ov, f'M{dr+1}:P{dr+DAYS}')

ch = LineChart(); ch.title = 'Running profit (units)'; ch.height = 10; ch.width = 20
ch.add_data(Reference(ov, min_col=16, min_row=dr, max_row=dr + DAYS), titles_from_data=True)
ch.set_categories(Reference(ov, min_col=12, min_row=dr + 1, max_row=dr + DAYS))
ch.legend = None; ch.x_axis.number_format = 'd mmm'; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.tickLblPos = 'low'
sr = ch.series[0]; sr.graphicalProperties.line.solidFill = GREEN; sr.graphicalProperties.line.width = 32000; sr.smooth = False
ov.add_chart(ch, 'L6')

ov.cell(r + 1, 1, 'theoverlay.com.au   18+   Gamble responsibly: 1800 858 858 or gamblinghelponline.org.au').font = font(size=9, color=MUTED)
wb.save('../Overlay results (live).xlsx')
print('saved')
